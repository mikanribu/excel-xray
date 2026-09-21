"""Bounded-memory portfolio scan and consolidated review exports.

The SQLite file is the single persisted assessment store. Original workbooks
remain in place; their bytes are only copied when a reviewer requests them.
"""

from __future__ import annotations

import csv
import hashlib
import heapq
import json
import os
import sqlite3
import sys
import time
import zipfile
from collections import defaultdict
from dataclasses import asdict
from pathlib import Path

from .assessment import assess, to_dict
from .corpus import (DUP_THRESHOLD, SIMILAR_THRESHOLD, SKELETON_DUP, Fingerprint,
                     _run_rationalization, fingerprint, rationalization_context,
                     similarity)
from .review_rules import shorten_external
from .scan import xray_workbook
from .tabular import FILE_FIELDS, TAB_FIELDS, fmt_value
from .util import safe_scan_message

DB_NAME = "portfolio.sqlite"
SUMMARY_NAME = "file_summary.csv"
EXCEL_NAME = "portfolio_review.xlsx"
SHEET_NAME = "worksheet_details.csv"
FINDINGS_NAME = "portfolio_findings.csv"
DIAGNOSTICS_NAME = "diagnostics.csv"

SUMMARY_COLUMNS = ["Record ID", "File ID", "File Name", "Scan Status", "Scan Error",
                   "No. of Sheets - Total", "No. of Sheets - Hidden",
                   "Cached Errors", "Tabs Requiring Review"] + [label for _, _, label in FILE_FIELDS
                                                               if label not in {"File ID", "File Name"}] + [
                   "Assessment Status", "Assessment Error"]
TAB_COLUMNS = ["Record ID", "File ID", "File Name", "Tab Name", "Original Position"] + [
    label for _, _, label in TAB_FIELDS if label != "Tab Name"]
DIAGNOSTIC_COLUMNS = ["Record ID", "File ID", "File Name", "Kind", "Sheet / Group",
                      "Error Type", "Count", "Area Type", "Affected Outputs", "Recalculation Needed",
                      "Reviewer Action", "Sample Cells"]
FINDING_COLUMNS = ["Scope", "Finding Type", "Severity", "Affected EUCs", "Affected Cells / Tabs",
                   "Finding", "Reviewer Action", "Evidence"]
MAX_MATCHES_PER_FILE = 20


def connect(run_dir: str | Path) -> sqlite3.Connection:
    path = Path(run_dir) / DB_NAME
    db = sqlite3.connect(path)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    db.executescript("""
        CREATE TABLE IF NOT EXISTS files (
          id INTEGER PRIMARY KEY, source_path TEXT NOT NULL UNIQUE, file_id TEXT,
          file_name TEXT NOT NULL, size_bytes INTEGER, source_mtime_ns INTEGER,
          sha256 TEXT, scan_status TEXT NOT NULL, scan_error TEXT,
          sheet_count INTEGER DEFAULT 0, hidden_count INTEGER DEFAULT 0,
          cached_error_count INTEGER DEFAULT 0, review_count INTEGER DEFAULT 0,
          summary_json TEXT, assessment_json TEXT, fingerprint_json TEXT
        );
        CREATE TABLE IF NOT EXISTS tabs (
          file_pk INTEGER NOT NULL REFERENCES files(id) ON DELETE CASCADE,
          position INTEGER NOT NULL, name TEXT NOT NULL, visibility TEXT,
          category TEXT, review_required TEXT, detail_json TEXT NOT NULL,
          PRIMARY KEY (file_pk, position)
        );
        CREATE TABLE IF NOT EXISTS errors (
          file_pk INTEGER NOT NULL REFERENCES files(id) ON DELETE CASCADE,
          sheet TEXT, error_type TEXT, count INTEGER, area_type TEXT,
          downstream_json TEXT, recalc_required INTEGER, reviewer_action TEXT,
          sample_json TEXT
        );
        CREATE TABLE IF NOT EXISTS hidden_groups (
          file_pk INTEGER NOT NULL REFERENCES files(id) ON DELETE CASCADE,
          purpose TEXT, count INTEGER, sheets_json TEXT,
          depended_json TEXT, explanation TEXT
        );
        CREATE TABLE IF NOT EXISTS pairs (
          file_a INTEGER NOT NULL REFERENCES files(id) ON DELETE CASCADE,
          file_b INTEGER NOT NULL REFERENCES files(id) ON DELETE CASCADE,
          score REAL NOT NULL, relationship TEXT NOT NULL, signals_json TEXT NOT NULL,
          PRIMARY KEY (file_a, file_b)
        );
        CREATE TABLE IF NOT EXISTS findings (
          id INTEGER PRIMARY KEY, scope TEXT, type TEXT, severity TEXT,
          affected_eucs INTEGER, affected_count INTEGER,
          title TEXT, action TEXT, evidence TEXT
        );
        CREATE INDEX IF NOT EXISTS files_status ON files(scan_status);
        CREATE INDEX IF NOT EXISTS tabs_file ON tabs(file_pk);
        CREATE INDEX IF NOT EXISTS errors_file ON errors(file_pk);
        CREATE INDEX IF NOT EXISTS pairs_b ON pairs(file_b);
    """)
    return db


def _json(value) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


def _safe_scan_error(exc: Exception, path: str) -> str:
    """Keep actionable scan status while removing machine-specific paths."""
    message = safe_scan_message(exc, path)
    return f"{getattr(exc, 'category', type(exc).__name__)}: {message}".strip()


def _display_scan_error(value) -> str:
    """Use a stable reviewer-facing value for successful scans without errors."""
    text = str(value).strip() if value is not None else ""
    return text or "N/A"


def _store_failure(db: sqlite3.Connection, path: str, exc: Exception) -> None:
    try:
        st = os.stat(path)
        size, mtime = st.st_size, st.st_mtime_ns
    except OSError:
        size, mtime = None, None
    db.execute("""INSERT INTO files(source_path,file_name,size_bytes,source_mtime_ns,scan_status,scan_error)
        VALUES(?,?,?,?,?,?) ON CONFLICT(source_path) DO UPDATE SET
        size_bytes=excluded.size_bytes,source_mtime_ns=excluded.source_mtime_ns,
        scan_status=excluded.scan_status,scan_error=excluded.scan_error,
        file_id=NULL,sha256=NULL,sheet_count=0,hidden_count=0,
        cached_error_count=0,review_count=0,
        summary_json=NULL,assessment_json=NULL,fingerprint_json=NULL""",
        (path, os.path.basename(path), size, mtime, "failed", _safe_scan_error(exc, path)))
    pk = db.execute("SELECT id FROM files WHERE source_path=?", (path,)).fetchone()[0]
    for table in ("tabs", "errors", "hidden_groups"):
        db.execute(f"DELETE FROM {table} WHERE file_pk=?", (pk,))


def _store_success(db: sqlite3.Connection, path: str, wx, a) -> None:
    st = os.stat(path)
    summary = {label: fmt_value(getattr(a.file, attr).value)
               for _, attr, label in FILE_FIELDS}
    summary["Assessment Status"] = a.scan.get("assessment_status") or "complete"
    summary["Assessment Error"] = a.scan.get("assessment_error") or "N/A"
    fp = fingerprint(wx, a)
    hidden = a.review.hidden_summary
    scan_error = None
    if wx.parse_status == "partial":
        details = "; ".join(wx.warnings[:3]) or "Workbook scan incomplete; review diagnostics"
        scan_error = "Partial: " + safe_scan_message(details, path)
    db.execute("""INSERT INTO files(source_path,file_id,file_name,size_bytes,source_mtime_ns,sha256,
        scan_status,scan_error,sheet_count,hidden_count,cached_error_count,review_count,
        summary_json,assessment_json,fingerprint_json)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(source_path) DO UPDATE SET
        file_id=excluded.file_id,file_name=excluded.file_name,size_bytes=excluded.size_bytes,
        source_mtime_ns=excluded.source_mtime_ns,sha256=excluded.sha256,
        scan_status=excluded.scan_status,scan_error=excluded.scan_error,sheet_count=excluded.sheet_count,
        hidden_count=excluded.hidden_count,cached_error_count=excluded.cached_error_count,
        review_count=excluded.review_count,summary_json=excluded.summary_json,
        assessment_json=excluded.assessment_json,fingerprint_json=excluded.fingerprint_json""",
        (path, a.file.file_id.value, wx.filename, st.st_size, st.st_mtime_ns,
         wx.sha256, wx.parse_status, scan_error, len(wx.sheets), hidden.get("hidden_count", 0),
         sum(g.count for g in a.review.error_groups),
         sum(t.human_validation_required.value == "Y" for t in a.tabs),
         _json(summary), _json(to_dict(a)), _json({
             **asdict(fp),
             "skeletons": sorted(fp.skeletons), "headers": sorted(fp.headers),
             "functions": sorted(fp.functions), "sheet_names": sorted(fp.sheet_names),
         })))
    pk = db.execute("SELECT id FROM files WHERE source_path=?", (path,)).fetchone()[0]
    for table in ("tabs", "errors", "hidden_groups"):
        db.execute(f"DELETE FROM {table} WHERE file_pk=?", (pk,))
    positions = {s.name: s.position for s in wx.sheets}
    for tab in a.tabs:
        db.execute("INSERT INTO tabs VALUES(?,?,?,?,?,?,?)",
                   (pk, positions.get(tab.tab_name.value, 0), tab.tab_name.value,
                    tab.tab_visibility.value, tab.tab_category.value,
                    tab.human_validation_required.value, _json(asdict(tab))))
    for g in a.review.error_groups:
        db.execute("INSERT INTO errors VALUES(?,?,?,?,?,?,?,?,?)",
                   (pk, g.sheet, g.error_type, g.count, g.area_type,
                    _json(g.downstream_outputs), int(g.recalc_required),
                    g.reviewer_action, _json(g.sample_cells)))
    for g in hidden.get("groups", []):
        if not isinstance(g, dict):
            g = asdict(g)
        db.execute("INSERT INTO hidden_groups VALUES(?,?,?,?,?,?)",
                   (pk, g["purpose"], g["count"], _json(g["sheets"]),
                    _json(g["depended_on_by"]), g["explanation"]))


def _restore_fp(raw: str) -> Fingerprint:
    d = json.loads(raw)
    for key in ("skeletons", "headers", "functions", "sheet_names"):
        d[key] = set(d[key])
    return Fingerprint(**d)


def _compare(db: sqlite3.Connection, assessor=None) -> None:
    """Compare compact fingerprints, without retaining workbook scans.

    Every pair is considered at 2,000 files (about two million lightweight
    comparisons). Only the strongest review matches per EUC are persisted,
    so an estate of identical templates cannot generate millions of rows.
    """
    rows = db.execute("SELECT id,fingerprint_json FROM files WHERE fingerprint_json IS NOT NULL ORDER BY id")
    fps = [(row["id"], _restore_fp(row["fingerprint_json"])) for row in rows]
    db.execute("DELETE FROM pairs")
    top: dict[int, list] = {pk: [] for pk, _ in fps}

    def offer(pk, peer, score):
        item = (score["overall"], peer, "Structural similarity review", _json(score))
        heap = top[pk]
        if len(heap) < MAX_MATCHES_PER_FILE:
            heapq.heappush(heap, item)
        elif item[:2] > heap[0][:2]:
            heapq.heapreplace(heap, item)

    for i, (a_id, a) in enumerate(fps):
        for b_id, b in fps[i + 1:]:
            score = similarity(a, b)
            if score["overall"] < SIMILAR_THRESHOLD and score["skeleton"] < SKELETON_DUP:
                continue
            offer(a_id, b_id, score)
            offer(b_id, a_id, score)
    review_pairs = {}
    for pk, heap in top.items():
        for score, peer, relationship, signals in heap:
            a_id, b_id = sorted((pk, peer))
            review_pairs[(a_id, b_id)] = (score, relationship, signals)
    db.executemany("INSERT INTO pairs VALUES(?,?,?,?,?)",
                   ((a, b, *details) for (a, b), details in review_pairs.items()))
    # Persist structural candidates as review shortlists only. A similarity
    # score is never written as a business duplication/consolidation finding.
    count = len(fps)
    matches = defaultdict(list)
    for p in db.execute("""SELECT p.*,a.file_name AS a_name,b.file_name AS b_name,
                                  a.file_id AS a_file_id,b.file_id AS b_file_id
                           FROM pairs p JOIN files a ON a.id=p.file_a JOIN files b ON b.id=p.file_b"""):
        signals = json.loads(p["signals_json"])
        for pk, peer, peer_id in ((p["file_a"], p["b_name"], p["b_file_id"]),
                                  (p["file_b"], p["a_name"], p["a_file_id"])):
            matches[pk].append({"file": shorten_external(peer), "file_id": peer_id,
                                "comparison_id": str(p["file_b"] if pk == p["file_a"] else p["file_a"]),
                                "similarity": p["score"], "signals": signals})

    stored = db.execute("SELECT id,summary_json,assessment_json,fingerprint_json FROM files "
                        "WHERE assessment_json IS NOT NULL").fetchall()
    rows_by_id = {r["id"]: r for r in stored}
    for row in stored:
        pk = row["id"]
        ms = sorted(matches[pk], key=lambda m: -m["similarity"])
        note = (f"{count - 1} other workbook(s) compared; structural matches are review "
                f"shortlists only; strongest {MAX_MATCHES_PER_FILE} retained per EUC")
        a = json.loads(row["assessment_json"])
        a.setdefault("file", {})
        a["file"].update({
            "potential_duplication": {
                "value": {"verdict": "Not established — functional comparison required",
                          "matches": ms, "corpus_compared": count - 1,
                          "similarity_threshold": DUP_THRESHOLD},
                "basis": "needs_corpus", "confidence": None,
                "evidence": [note, "shared process labels or structural similarity alone do not establish duplicate activity"],
            },
            "similar_duplicate_files": {
                "value": "Not established — functional comparison required",
                "basis": "needs_corpus", "confidence": None, "evidence": [note],
            },
            "potential_consolidation": {
                "value": {"verdict": "Not established — functional comparison required",
                          "candidates": [], "structural_candidates": ms},
                "basis": "needs_corpus", "confidence": None,
                "evidence": [note, "consolidation requires business-purpose fit and blocker review"],
            },
        })

        if callable(getattr(assessor, "rationalize", None)) and row["fingerprint_json"]:
            try:
                assessment = assessment_from_json(_json(a))
                fp = _restore_fp(row["fingerprint_json"])
                current = rationalization_context(a, fp, comparison_id=str(pk))
                current["corpus_compared"] = count - 1
                peers = []
                for match in ms:
                    peer_row = rows_by_id.get(int(match["comparison_id"]))
                    if not peer_row or not peer_row["fingerprint_json"]:
                        continue
                    peer_assessment = json.loads(peer_row["assessment_json"])
                    peer_fp = _restore_fp(peer_row["fingerprint_json"])
                    peer = rationalization_context(
                        peer_assessment, peer_fp, comparison_id=str(peer_row["id"]))
                    peer["structural_similarity"] = match["signals"]
                    peers.append(peer)
                _run_rationalization(assessment, current, peers, assessor)
                from .assessment import to_dict
                a = to_dict(assessment)
            except (KeyError, TypeError, ValueError):
                # Older/partial assessment JSON can still retain structural
                # review candidates even when it cannot enter the LLM stage.
                pass

        summary = json.loads(row["summary_json"]) if row["summary_json"] else {}
        for _, attr, label in FILE_FIELDS:
            field_value = (a.get("file", {}).get(attr) or {}).get("value")
            summary[label] = fmt_value(field_value)
        db.execute("UPDATE files SET assessment_json=?,summary_json=? WHERE id=?",
                   (_json(a), _json(summary), pk))


def _findings(db: sqlite3.Connection) -> None:
    db.execute("DELETE FROM findings")
    def add(kind, severity, files, cells, title, action, evidence):
        db.execute("INSERT INTO findings(scope,type,severity,affected_eucs,affected_count,title,action,evidence) VALUES(?,?,?,?,?,?,?,?)",
                   ("Portfolio", kind, severity, files, cells, title, action, evidence))

    for r in db.execute("""SELECT error_type,area_type,COUNT(DISTINCT file_pk) AS files,
                           SUM(count) AS cells FROM errors GROUP BY error_type,area_type
                           ORDER BY cells DESC"""):
        severity = "High" if r["area_type"] == "final output" or r["error_type"] == "#REF!" else "Review"
        outputs = set()
        for detail in db.execute("SELECT downstream_json FROM errors WHERE error_type=? AND area_type=?",
                                 (r["error_type"], r["area_type"])):
            outputs.update(json.loads(detail["downstream_json"]))
        output_note = (" Affected output tabs: " + ", ".join(sorted(outputs)[:8])
                       + (" and others." if len(outputs) > 8 else ".")) if outputs else " No downstream output identified."
        action = ("Repair the broken reference; recalculation alone will not fix it."
                  if r["error_type"] == "#REF!" else
                  "Refresh linked sources and recalculate, then confirm whether errors persist.")
        add("Cached errors", severity, r["files"], r["cells"],
            f"{r['error_type']} in {r['area_type']}",
            action + " Check affected business outputs.",
            f"{r['files']} workbook(s), {r['cells']} cached error cell(s)." + output_note)
    for r in db.execute("""SELECT purpose,COUNT(DISTINCT file_pk) AS files,SUM(count) AS tabs
                           FROM hidden_groups GROUP BY purpose ORDER BY tabs DESC"""):
        add("Hidden sheets", "Review", r["files"], r["tabs"],
            f"Hidden {r['purpose'].lower()} tabs", "Confirm the role and downstream use before consolidating or retiring a workbook.",
            f"{r['tabs']} tab(s) in {r['files']} workbook(s).")
    for r in db.execute("SELECT relationship,COUNT(*) AS pairs FROM pairs GROUP BY relationship"):
        files = db.execute("SELECT COUNT(DISTINCT id) FROM files WHERE id IN (SELECT file_a FROM pairs WHERE relationship=? UNION SELECT file_b FROM pairs WHERE relationship=?)",
                           (r["relationship"], r["relationship"])).fetchone()[0]
        add("Cross-EUC comparison", "Review",
            files, r["pairs"], r["relationship"],
            "Compare business purpose, inputs, calculations, controls and outputs before deciding whether the EUCs are similar or could be consolidated.",
            f"{r['pairs']} structural review pair(s); up to {MAX_MATCHES_PER_FILE} strongest matches per EUC. "
            "Similarity scores do not establish functional duplication.")
    r = db.execute("SELECT COUNT(*) AS files,SUM(review_count) AS tabs FROM files WHERE review_count>0").fetchone()
    if r["files"]:
        add("Human validation", "Review", r["files"], r["tabs"],
            "Tabs needing human validation", "Use the worksheet details and each tab's validation reason to assign review.",
            f"{r['tabs']} tab(s) across {r['files']} EUC(s).")
    rationalization = {
        key: {"files": set(), "items": set(), "examples": []}
        for key in ("duplication", "similarity", "simplification", "consolidation",
                    "automation", "retirement")
    }
    file_names = {r["id"]: r["file_name"] for r in
                  db.execute("SELECT id,file_name FROM files")}

    def value_of(file_fields, name):
        fld = file_fields.get(name) or {}
        return fld.get("value") if isinstance(fld, dict) else None

    def record_pair(kind, current_id, candidate):
        try:
            peer_id = int(candidate.get("comparison_id"))
        except (TypeError, ValueError):
            return
        if peer_id not in file_names or current_id == peer_id:
            return
        pair = tuple(sorted((current_id, peer_id)))
        group = rationalization[kind]
        group["items"].add(pair)
        group["files"].update(pair)
        if len(group["examples"]) < 5:
            example = f"{shorten_external(file_names[current_id])} ↔ {shorten_external(file_names[peer_id])}"
            if example not in group["examples"]:
                group["examples"].append(example)

    def record_local(kind, file_id, candidates):
        group = rationalization[kind]
        for index, candidate in enumerate(candidates):
            if not isinstance(candidate, dict):
                continue
            group["files"].add(file_id)
            group["items"].add((file_id, index))
            if len(group["examples"]) < 5:
                label = candidate.get("worksheet") or candidate.get("observed_manual_step") or "workbook"
                group["examples"].append(f"{shorten_external(file_names[file_id])}: {label}")

    manual_cues = ("manual", "paste", "override", "keyed", "hardcode", "hard-coded")
    for row in db.execute("SELECT id,file_name,assessment_json FROM files WHERE assessment_json IS NOT NULL"):
        current_id = row["id"]
        file_fields = json.loads(row["assessment_json"]).get("file", {})
        duplication = value_of(file_fields, "potential_duplication") or {}
        if (isinstance(duplication, dict)
                and str(duplication.get("verdict", "")).startswith("Potential functional duplication")):
            for candidate in duplication.get("matches", []):
                if isinstance(candidate, dict):
                    record_pair("duplication", current_id, candidate)
        similar = value_of(file_fields, "similar_duplicate_files") or []
        for candidate in similar if isinstance(similar, list) else []:
            if isinstance(candidate, dict):
                record_pair("similarity", current_id, candidate)
        consolidation = value_of(file_fields, "potential_consolidation") or {}
        if (isinstance(consolidation, dict)
                and str(consolidation.get("verdict", "")).startswith("Potential consolidation")):
            for candidate in consolidation.get("candidates", []):
                if isinstance(candidate, dict):
                    record_pair("consolidation", current_id, candidate)

        simplification = value_of(file_fields, "potential_simplification") or {}
        local_simplification = (simplification.get("candidates", [])
                                if isinstance(simplification, dict) else [])
        record_local("simplification", current_id, local_simplification)

        automation = value_of(file_fields, "potential_automation") or {}
        local_automation = (automation.get("candidates", [])
                            if isinstance(automation, dict) else [])
        local_automation = [candidate for candidate in local_automation
                            if isinstance(candidate, dict) and any(
                                any(cue in str(text).casefold() for cue in manual_cues)
                                for text in ([candidate.get("suspected_manual_step"),
                                              candidate.get("observed_manual_step")]
                                             + [e.get("text", "") if isinstance(e, dict) else e
                                                for e in candidate.get("evidence", [])])
                            )]
        record_local("automation", current_id, local_automation)

        retirement = value_of(file_fields, "potential_retirement") or {}
        if (isinstance(retirement, dict)
                and str(retirement.get("verdict", "")).startswith("Potential retirement")):
            record_local("retirement", current_id, [retirement])

    rollup_specs = (
        ("duplication", "Potential functional duplication",
         "Review the cited business similarities and differences with process owners."),
        ("similarity", "Similar EUC business use cases",
         "Compare the listed purposes, outputs and material differences."),
        ("simplification", "Potential simplification",
         "Review the evidence-backed changes with process owners and confirm the business benefit."),
        ("consolidation", "Potential consolidation",
         "Review the shared-solution proposal and documented blockers with process owners."),
        ("automation", "Potential automation",
         "Confirm the evidenced manual step, repeatability, inputs, exceptions and controls."),
        ("retirement", "Potential retirement",
         "Confirm usage, ownership, recipient needs, replacement coverage and retention obligations."),
    )
    for key, title, action in rollup_specs:
        group = rationalization[key]
        if group["items"]:
            examples = "; examples: " + "; ".join(group["examples"]) if group["examples"] else ""
            add("Cross-EUC rationalisation", "Review", len(group["files"]), len(group["items"]),
                title, action,
                f"{len(group['items'])} evidence-backed candidate(s) across {len(group['files'])} EUC(s)."
                + examples)
    failed = db.execute("SELECT COUNT(*) FROM files WHERE scan_status='failed'").fetchone()[0]
    if failed:
        add("Coverage", "High", failed, 0, "Unreadable workbooks",
            "Resolve the scan errors and resume the run.", f"{failed} submitted EUC(s) could not be assessed.")


def scan_portfolio(paths: list[str], run_dir: str | Path, *, max_rows: int = 200_000,
                   assessor=None, resume: bool = False) -> dict:
    run_dir = Path(run_dir)
    run_dir.mkdir(parents=True, exist_ok=True)
    db = connect(run_dir)
    done = failed = skipped = 0
    try:
        expected = {os.path.abspath(p) for p in paths}
        with db:
            for row in db.execute("SELECT id,source_path FROM files").fetchall():
                if row["source_path"] not in expected:
                    db.execute("DELETE FROM files WHERE id=?", (row["id"],))
        for n, item in enumerate(paths, 1):
            path = os.path.abspath(item)
            start = time.perf_counter()
            try:
                st = os.stat(path)
                previous = db.execute("SELECT * FROM files WHERE source_path=?", (path,)).fetchone()
                if (resume and previous and previous["scan_status"] != "failed"
                    and (previous["size_bytes"], previous["source_mtime_ns"])
                    == (st.st_size, st.st_mtime_ns) and source_is_current(previous)):
                    skipped += 1
                    continue
                wx = xray_workbook(path, max_rows=max_rows)
                a = assess(wx, assessor)
                with db:
                    _store_success(db, path, wx, a)
                done += 1
                print(f"[{n}/{len(paths)}] {wx.parse_status}: {wx.filename} ({time.perf_counter()-start:.1f}s)", file=sys.stderr)
            except Exception as exc:  # keep the rest of the batch usable
                with db:
                    _store_failure(db, path, exc)
                failed += 1
                print(f"[{n}/{len(paths)}] FAILED: {os.path.basename(path)}: {exc}", file=sys.stderr)
        with db:
            _compare(db, assessor)
            _findings(db)
        export_all(db, run_dir)
        return {"scanned": done, "failed": failed, "reused": skipped,
                "stored": db.execute("SELECT COUNT(*) FROM files").fetchone()[0],
                "run_dir": str(run_dir)}
    finally:
        db.close()


def selected_files(db: sqlite3.Connection, ids: list[int] | None = None):
    if ids is None:
        yield from db.execute("SELECT * FROM files ORDER BY file_name,id")
    else:
        for pk in dict.fromkeys(ids):
            row = db.execute("SELECT * FROM files WHERE id=?", (pk,)).fetchone()
            if row is not None:
                yield row


def _safe(value):
    """Prevent spreadsheet formula execution from workbook-supplied text."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def summary_rows(db, ids=None):
    for r in selected_files(db, ids):
        summary = json.loads(r["summary_json"]) if r["summary_json"] else {}
        failed = r["scan_status"] == "failed"
        not_assessed = "Not assessed — workbook scan failed"
        values = [r["id"], r["file_id"] or ("Not available — scan failed" if failed else "Not recorded"),
                  r["file_name"], r["scan_status"], _display_scan_error(r["scan_error"]),
                  "Not assessed" if failed else r["sheet_count"],
                  "Not assessed" if failed else r["hidden_count"],
                  "Not assessed" if failed else r["cached_error_count"],
                  "Not assessed" if failed else r["review_count"]]
        file_fields = [summary.get(c, not_assessed if failed else "Not recorded — rescan to populate")
                       for c in SUMMARY_COLUMNS[9:-2]]
        assessment_status = ("Not assessed" if failed else
                             summary.get("Assessment Status", "Not recorded — rescan to populate"))
        assessment_error = ("Workbook could not be scanned; see Scan Error." if failed else
                            summary.get("Assessment Error", "N/A"))
        yield [_safe(v) for v in values + file_fields + [assessment_status, assessment_error]]


def tab_detail_rows(db, ids=None):
    for f in selected_files(db, ids):
        for r in db.execute("SELECT * FROM tabs WHERE file_pk=? ORDER BY visibility!='Visible',position", (f["id"],)):
            detail = json.loads(r["detail_json"])
            values = [f["id"], f["file_id"], f["file_name"], r["name"], r["position"] + 1]
            values += [fmt_value(detail[attr]["value"]) for _, attr, _ in TAB_FIELDS if attr != "tab_name"]
            yield [_safe(v) for v in values]


def diagnostic_rows(db, ids=None):
    for f in selected_files(db, ids):
        base = [f["id"], f["file_id"] or ("Not available — scan failed"
                                            if f["scan_status"] == "failed" else "Not recorded"),
                f["file_name"]]
        if f["scan_status"] == "failed":
            yield [_safe(v) for v in base + ["Scan failure", "Not applicable", "Workbook scan failed", "Not assessed",
                "Not assessed", "Not assessed", "Not assessed", "Resolve the Scan Error and rescan the workbook.",
                _display_scan_error(f["scan_error"])]]
        for r in db.execute("SELECT * FROM errors WHERE file_pk=? ORDER BY count DESC", (f["id"],)):
            yield [_safe(v) for v in base + ["Cached error", r["sheet"], r["error_type"], r["count"],
                r["area_type"], "; ".join(json.loads(r["downstream_json"])),
                "Yes" if r["recalc_required"] else "No", r["reviewer_action"],
                "; ".join(json.loads(r["sample_json"]))]]
        for r in db.execute("SELECT * FROM hidden_groups WHERE file_pk=? ORDER BY count DESC", (f["id"],)):
            yield [_safe(v) for v in base + ["Hidden sheets", r["purpose"], "", r["count"], "",
                "; ".join(json.loads(r["depended_json"])), "", r["explanation"],
                "; ".join(json.loads(r["sheets_json"]))]]


def finding_rows(db, ids=None):
    if ids is None:
        for r in db.execute("SELECT * FROM findings ORDER BY CASE severity WHEN 'High' THEN 0 ELSE 1 END,affected_eucs DESC"):
            yield [_safe(r[c]) for c in ("scope", "type", "severity", "affected_eucs",
                                          "affected_count", "title", "action", "evidence")]
        return
    chosen = {r["id"] for r in selected_files(db, ids)}
    errors = defaultdict(lambda: [set(), 0])
    hidden = defaultdict(lambda: [set(), 0])
    reviews = [0, 0]
    failed = 0
    for f in selected_files(db, ids):
        if f["scan_status"] == "failed":
            failed += 1
        if f["review_count"]:
            reviews[0] += 1
            reviews[1] += f["review_count"]
        for r in db.execute("SELECT error_type,area_type,count FROM errors WHERE file_pk=?", (f["id"],)):
            group = errors[(r["error_type"], r["area_type"])]
            group[0].add(f["id"])
            group[1] += r["count"]
        for r in db.execute("SELECT purpose,count FROM hidden_groups WHERE file_pk=?", (f["id"],)):
            group = hidden[r["purpose"]]
            group[0].add(f["id"])
            group[1] += r["count"]
    for (error_type, area), (files, cells) in sorted(errors.items(), key=lambda x: -x[1][1]):
        yield ["Selected EUCs", "Cached errors", "High" if area == "final output" or error_type == "#REF!" else "Review",
               len(files), cells, f"{error_type} in {area}",
               ("Repair the broken reference; recalculation alone will not fix it."
                if error_type == "#REF!" else "Refresh sources and recalculate, then check affected outputs."),
               f"{len(files)} workbook(s), {cells} cached error cell(s)."]
    for purpose, (files, count) in sorted(hidden.items(), key=lambda x: -x[1][1]):
        yield ["Selected EUCs", "Hidden sheets", "Review", len(files), count,
               f"Hidden {purpose.lower()} tabs",
               "Confirm role and downstream use before consolidating or retiring a workbook.",
               f"{count} tabs in {len(files)} workbook(s)."]
    relationships = defaultdict(lambda: [set(), 0])
    for r in db.execute("SELECT file_a,file_b,relationship FROM pairs"):
        if r["file_a"] in chosen and r["file_b"] in chosen:
            group = relationships[r["relationship"]]
            group[0].update((r["file_a"], r["file_b"]))
            group[1] += 1
    for relationship, (files, count) in relationships.items():
        yield ["Selected EUCs", "Cross-EUC comparison", "Review",
               len(files), count, relationship,
               "Confirm business purpose and outputs before any consolidation.",
               f"{count} pair(s) within the selected EUCs."]
    if reviews[0]:
        yield ["Selected EUCs", "Human validation", "Review", *reviews,
               "Tabs needing human validation", "Use each tab's validation reason to assign review.",
               f"{reviews[1]} tabs across {reviews[0]} EUCs."]
    opportunities = {"potential_simplification": [0, 0], "potential_automation": [0, 0]}
    for f in selected_files(db, ids):
        if not f["assessment_json"]:
            continue
        file_fields = json.loads(f["assessment_json"])["file"]
        for key in opportunities:
            value = file_fields.get(key, {}).get("value") or {}
            candidates = value.get("candidates", []) if isinstance(value, dict) else []
            if key == "potential_automation":
                candidates = [c for c in candidates if str(c.get("suspected_manual_step", ""))
                              .lower().startswith("explicit manual-entry/override area")]
            if candidates:
                opportunities[key][0] += 1
                opportunities[key][1] += len(candidates)
    for key, title, action in (
        ("potential_simplification", "Simplification candidates",
         "Confirm business benefit with process owners."),
        ("potential_automation", "Manual-workflow review leads",
         "Confirm manual steps, frequency, inputs, exceptions and approval controls before assessing automation suitability."),
    ):
        files, tabs = opportunities[key]
        if files:
            yield ["Selected EUCs", "Cross-EUC opportunities", "Review", files, tabs,
                   title, action, f"{tabs} candidate tab(s) across {files} workbook(s)."]
    if failed:
        yield ["Selected EUCs", "Coverage", "High", failed, 0, "Unreadable workbooks",
               "Resolve scan errors and resume the run.", f"{failed} EUC(s) could not be assessed."]


def write_csv(path, headers, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(rows)
    return path


def write_excel(path, db, ids=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    for name, headers, rows in (
        ("File summary", SUMMARY_COLUMNS, summary_rows(db, ids)),
        ("Worksheet details", TAB_COLUMNS, tab_detail_rows(db, ids)),
        ("Diagnostics", DIAGNOSTIC_COLUMNS, diagnostic_rows(db, ids)),
        ("Portfolio findings", FINDING_COLUMNS, finding_rows(db, ids)),
    ):
        ws = wb.create_sheet(name)
        ws.freeze_panes = "A2"
        for col, width in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(col)].width = min(54, max(16, len(width) + 3))
        header_cells = []
        for h in headers:
            cell = WriteOnlyCell(ws, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24558B")
            header_cells.append(cell)
        ws.append(header_cells)
        count = 1
        for row in rows:
            ws.append(row)
            count += 1
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"
    wb.save(path)
    return path


def export_all(db, run_dir):
    run_dir = Path(run_dir)
    write_csv(run_dir / SUMMARY_NAME, SUMMARY_COLUMNS, summary_rows(db))
    write_csv(run_dir / SHEET_NAME, TAB_COLUMNS, tab_detail_rows(db))
    write_csv(run_dir / DIAGNOSTICS_NAME, DIAGNOSTIC_COLUMNS, diagnostic_rows(db))
    write_csv(run_dir / FINDINGS_NAME, FINDING_COLUMNS, finding_rows(db))
    write_excel(run_dir / EXCEL_NAME, db)


def bundle(path, db, ids):
    """Selected originals plus their review files; zip members are unambiguous."""
    from tempfile import TemporaryDirectory
    with TemporaryDirectory() as tmp:
        csv_path = Path(tmp) / SUMMARY_NAME
        xlsx_path = Path(tmp) / EXCEL_NAME
        findings_path = Path(tmp) / FINDINGS_NAME
        write_csv(csv_path, SUMMARY_COLUMNS, summary_rows(db, ids))
        write_excel(xlsx_path, db, ids)
        write_csv(findings_path, FINDING_COLUMNS, finding_rows(db, ids))
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as zf:
            zf.write(csv_path, SUMMARY_NAME)
            zf.write(xlsx_path, EXCEL_NAME)
            zf.write(findings_path, FINDINGS_NAME)
            for r in selected_files(db, ids):
                original = Path(r["source_path"])
                if not source_is_current(r):
                    raise FileNotFoundError(f"Original missing or changed since scan: {original}")
                zf.write(original, f"originals/{r['id']}_{original.name}")
    return path


def source_is_current(row) -> bool:
    path = Path(row["source_path"])
    if not path.is_file():
        return False
    stat = path.stat()
    if stat.st_size != row["size_bytes"]:
        return False
    if row["sha256"]:
        h = hashlib.sha256()
        with path.open("rb") as fh:
            for block in iter(lambda: fh.read(1024 * 1024), b""):
                h.update(block)
        return h.hexdigest() == row["sha256"]
    return stat.st_mtime_ns == row["source_mtime_ns"]


def assessment_from_json(raw: str):
    """Restore the stored assessment for the existing individual report code."""
    from .assessment import Assessment, Field, FileAssessment, ReviewFindings, TabAssessment
    from .review_rules import ErrorGroup, HiddenGroup
    d = json.loads(raw)
    file_assessment = FileAssessment(**{k: Field(**v) for k, v in d["file"].items()})
    tabs = [TabAssessment(**{k: Field(**v) for k, v in row.items()}) for row in d["tabs"]]
    review = d["review"]
    review["error_groups"] = [ErrorGroup(**g) for g in review["error_groups"]]
    review["hidden_summary"]["groups"] = [HiddenGroup(**g)
                                           for g in review["hidden_summary"].get("groups", [])]
    return Assessment(file=file_assessment, tabs=tabs, review=ReviewFindings(**review),
                      scan=d.get("scan", {}))
