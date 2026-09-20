"""Excel (.xlsx) report exporter — the default report format (Step 21)."""

from __future__ import annotations

import subprocess
import sys

import openpyxl

from excel_xray import assess
from excel_xray.xlsx_report import build_xlsx_report

REQUIRED_SHEETS = {
    "File assessment", "Tab assessments", "Sheet inventory", "Regions",
    "Formula patterns", "Warnings", "Report information", "Error summary",
    "Hidden sheet groups", "Input sources", "Calculation steps",
    "Review opportunities",
}


def test_xlsx_report_has_every_required_sheet(xray):
    wb = build_xlsx_report(xray, assess(xray))
    assert REQUIRED_SHEETS <= set(wb.sheetnames)


def test_assessment_sheets_carry_reviewer_columns(xray):
    wb = build_xlsx_report(xray, assess(xray))
    for name in ("File assessment", "Tab assessments", "Review opportunities"):
        headers = [c.value for c in next(wb[name].iter_rows(min_row=1, max_row=1))]
        assert "Reviewer Value" in headers
        assert "Reviewer Notes" in headers


def test_assessment_sheets_are_filterable_and_frozen(xray):
    wb = build_xlsx_report(xray, assess(xray))
    ws = wb["File assessment"]
    assert ws.auto_filter.ref is not None
    assert ws.freeze_panes == "A2"


def test_reviewer_export_has_business_summary_and_keeps_formula_details_in_appendix(xray):
    wb = build_xlsx_report(xray, assess(xray))
    file_fields = [row[1].value for row in wb["File assessment"].iter_rows(min_row=2)]
    assert "Logic Types" in file_fields
    assert "Logic Type" not in file_fields
    info = " ".join(str(c.value) for row in wb["Report information"].iter_rows() for c in row)
    assert "Scan status" in info and "No. of Sheets - Total" in info
    assert "SHA256" not in info and "Content hash" not in info and "Size (bytes)" not in info
    calc_headers = [c.value for c in next(wb["Calculation steps"].iter_rows(min_row=1, max_row=1))]
    assert calc_headers == ["Worksheet", "Business calculation or transformation"]
    input_headers = [c.value for c in next(wb["Input sources"].iter_rows(min_row=1, max_row=1))]
    assert input_headers == ["Business purpose", "Source type", "Source name", "Reference count",
                             "Consuming worksheets", "Source status", "Essentiality"]
    all_values = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row)
    assert xray.path not in all_values
    assert "full_path_evidence" not in all_values
    assert "Top formula shapes" in [c.value for c in next(wb["Formula patterns"].iter_rows(min_row=1, max_row=1))]


def test_error_summary_sheet_has_grouped_and_detail_blocks(xray):
    wb = build_xlsx_report(xray, assess(xray))
    ws = wb["Error summary"]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[0][0] == "Error type"
    assert any(row[0] == "Detail: individual cell references" for row in values)


def test_hidden_sheet_groups_sheet_reads_x_of_y(xray):
    wb = build_xlsx_report(xray, assess(xray))
    ws = wb["Hidden sheet groups"]
    header_cell = ws.cell(row=1, column=1).value
    assert "worksheets are hidden" in header_cell


def test_estate_sheet_only_present_when_estate_supplied(xray):
    wb = build_xlsx_report(xray, assess(xray))
    assert "Estate comparison" not in wb.sheetnames


def test_cli_defaults_to_xlsx(tmp_path, fixture_path):
    out = tmp_path / "out"
    out.mkdir()
    subprocess.run(
        [sys.executable, "-m", "excel_xray", fixture_path, "-o", str(out)],
        check=True, capture_output=True, text=True,
    )
    files = list(out.glob("**/*.xlsx"))
    assert files, "default format must be .xlsx"
    wb = openpyxl.load_workbook(files[0])
    assert "File assessment" in wb.sheetnames


def test_cli_html_format_still_available(tmp_path, fixture_path):
    out = tmp_path / "out"
    out.mkdir()
    subprocess.run(
        [sys.executable, "-m", "excel_xray", fixture_path, "-o", str(out), "--format", "html"],
        check=True, capture_output=True, text=True,
    )
    files = list(out.glob("**/*.html"))
    assert files, "--format html must still produce an HTML report"
    html = files[0].read_text()
    assert "EUC assessment" in html
