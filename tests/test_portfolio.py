"""Portfolio output is consolidated, selectable and faithful to source files."""

from __future__ import annotations

import csv
import hashlib
import json
import zipfile
from pathlib import Path

import openpyxl

from excel_xray.portfolio import (DB_NAME, EXCEL_NAME, SUMMARY_NAME, SUMMARY_COLUMNS,
                                  _compare, _safe_scan_error, _store_success, summary_rows,
                                  assessment_from_json, bundle, connect,
                                  scan_portfolio, source_is_current, write_excel)
from excel_xray import assess
from excel_xray.report import build_report
from excel_xray.scan import xray_workbook


FIXTURES = Path(__file__).parent / "fixtures"


def test_portfolio_one_file_row_per_euc_and_tab_drilldown(tmp_path):
    paths = [str(FIXTURES / "complex_euc_model.xlsx"),
             str(FIXTURES / "messy_reserving_model.xlsx")]
    result = scan_portfolio(paths, tmp_path)
    assert result["stored"] == 2
    assert not list(tmp_path.glob("*.json"))
    assert not list(tmp_path.glob("xray_*.xlsx"))
    with (tmp_path / SUMMARY_NAME).open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 2
    assert {r["File Name"] for r in rows} == {Path(p).name for p in paths}
    assert all(r["File ID"] and r["Scan Status"] == "full" for r in rows)
    assert "No. of Sheets - Total" in rows[0]
    assert "No. of Sheets - Hidden" in rows[0]
    assert rows[0]["Scan Error"] == "N/A"
    assert "Business Area Purpose" in rows[0]
    assert "Business Area Process" not in rows[0]
    assert not {"Source Path", "Size Bytes", "SHA256"} & set(rows[0])
    assert not any(str(tmp_path) in str(value) for row in rows for value in row.values())
    assert "Process" in rows[0] and "Sub-Process" in rows[0]
    wb = openpyxl.load_workbook(tmp_path / EXCEL_NAME, read_only=True)
    assert wb.sheetnames == ["File summary", "Worksheet details", "Diagnostics", "Portfolio findings"]
    file_summary_rows = list(wb["File summary"].values)
    assert len(file_summary_rows) == 3
    summary_headers = set(file_summary_rows[0])
    assert "No. of Sheets - Total" in summary_headers
    assert "No. of Sheets - Hidden" in summary_headers
    assert {"Business Area Purpose", "Process", "Sub-Process", "Logic Types"} <= summary_headers
    assert "Business Area Process" not in summary_headers
    assert not {"Source Path", "Size Bytes", "SHA256", "Logic Type"} & summary_headers
    assert not any(str(tmp_path) in str(value) for row in file_summary_rows for value in row)
    assert sum(1 for _ in wb["Worksheet details"].values) == 16
    with connect(tmp_path) as db:
        assert db.execute("SELECT COUNT(*) FROM tabs").fetchone()[0] == 15
        row = db.execute("SELECT * FROM files ORDER BY id LIMIT 1").fetchone()
        assert source_is_current(row)
        a = assessment_from_json(row["assessment_json"])
        html = build_report(xray_workbook(row["source_path"]), a)
        assert "EUC assessment" in html
        from excel_xray.xlsx_report import write_xlsx_report
        individual = tmp_path / "individual.xlsx"
        write_xlsx_report(xray_workbook(row["source_path"]), a, str(individual))
        individual_wb = openpyxl.load_workbook(individual, read_only=True)
        info = list(individual_wb["Report information"].values)
        assert ("Scan error", "N/A") in info
        individual_fields = [r[0] for r in individual_wb["File assessment"].iter_rows(
            min_row=2, min_col=2, max_col=2, values_only=True)]
        assert "Business Area Purpose" in individual_fields
        assert "Business Area Process" not in individual_fields


def test_failed_scan_error_is_concise_and_does_not_expose_path_or_credentials(tmp_path):
    path = str(tmp_path / "private" / "corrupt.xlsx")
    message = _safe_scan_error(ValueError(f"cannot read {path}; password=secret"), path)
    assert "corrupt.xlsx" in message
    assert path not in message
    assert "secret" not in message
    assert "password=[redacted]" in message
    assert "Source Path" not in SUMMARY_COLUMNS
    assert "SHA256" not in SUMMARY_COLUMNS


def test_partial_scan_reason_is_in_summary_without_source_path(tmp_path):
    source = str(FIXTURES / "messy_reserving_model.xlsx")
    wx = xray_workbook(source)
    wx.parse_status = "partial"
    wx.warnings.append(f"Could not finish scanning {source}")
    a = assess(wx)
    with connect(tmp_path) as db:
        _store_success(db, source, wx, a)
        row = next(summary_rows(db))
        assert row[3] == "partial"
        assert "Could not finish scanning" in row[4]
        assert source not in row[4]
        assert "partial" in a.scan["status"]
        assert source not in (a.scan["error"] or "")


def test_successful_scan_error_is_na_and_real_error_is_retained(tmp_path):
    from excel_xray.portfolio import _display_scan_error, _store_failure

    assert _display_scan_error(None) == "N/A"
    assert _display_scan_error("") == "N/A"
    assert _display_scan_error("#REF! could not be read") == "#REF! could not be read"
    corrupt = tmp_path / "broken.xlsx"
    corrupt.write_bytes(b"not an OOXML workbook")
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    with connect(run_dir) as db:
        _store_failure(db, str(corrupt), ValueError("invalid workbook format"))
        row = next(summary_rows(db))
        assert row[4] == "ValueError: invalid workbook format"


def test_selected_bundle_contains_only_selected_original_and_analysis(tmp_path):
    paths = [str(FIXTURES / "complex_euc_model.xlsx"),
             str(FIXTURES / "messy_reserving_model.xlsx")]
    scan_portfolio(paths, tmp_path)
    with connect(tmp_path) as db:
        row = db.execute("SELECT * FROM files ORDER BY id LIMIT 1").fetchone()
        dest = tmp_path / "selected.zip"
        bundle(dest, db, [row["id"]])
        with zipfile.ZipFile(dest) as zf:
            originals = [n for n in zf.namelist() if n.startswith("originals/")]
            assert len(originals) == 1
            assert hashlib.sha256(zf.read(originals[0])).hexdigest() == row["sha256"]
            with zf.open(SUMMARY_NAME) as fh:
                assert len(list(csv.reader(line.decode("utf-8-sig") for line in fh))) == 2
            with zf.open(EXCEL_NAME) as fh:
                assert fh.read(2) == b"PK"
        selected_excel = tmp_path / "selected.xlsx"
        write_excel(selected_excel, db, [row["id"]])
        wb = openpyxl.load_workbook(selected_excel, read_only=True)
        assert sum(1 for _ in wb["File summary"].values) == 2


def test_resume_reuses_unchanged_workbooks_and_failed_rows_are_retained(tmp_path):
    good = str(FIXTURES / "messy_reserving_model.xlsx")
    bad = tmp_path / "corrupt.xlsx"
    bad.write_bytes(b"not an OOXML workbook")
    first = scan_portfolio([good, str(bad)], tmp_path / "run")
    assert first["stored"] == 2 and first["failed"] == 1
    second = scan_portfolio([good, str(bad)], tmp_path / "run", resume=True)
    assert second["reused"] == 1 and second["failed"] == 1
    with connect(tmp_path / "run") as db:
        assert db.execute("SELECT COUNT(*) FROM files").fetchone()[0] == 2
        assert db.execute("SELECT scan_error FROM files WHERE scan_status='failed'").fetchone()[0]


def test_source_change_blocks_original_export(tmp_path):
    source = tmp_path / "source.xlsx"
    source.write_bytes((FIXTURES / "messy_reserving_model.xlsx").read_bytes())
    scan_portfolio([str(source)], tmp_path / "run")
    with connect(tmp_path / "run") as db:
        row = db.execute("SELECT * FROM files").fetchone()
        assert source_is_current(row)
        source.write_bytes(source.read_bytes() + b"changed")
        assert not source_is_current(row)
        try:
            bundle(tmp_path / "changed.zip", db, [row["id"]])
        except FileNotFoundError:
            pass
        else:
            assert False, "bundle should reject changed originals"


def test_identical_workbooks_are_structural_candidates_not_duplicate_findings(tmp_path):
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    content = (FIXTURES / "messy_reserving_model.xlsx").read_bytes()
    first.write_bytes(content)
    second.write_bytes(content)
    scan_portfolio([str(first), str(second)], tmp_path / "run")
    with connect(tmp_path / "run") as db:
        for row in db.execute("SELECT assessment_json,summary_json FROM files"):
            assessment = json.loads(row["assessment_json"])
            value = assessment["file"]["potential_duplication"]["value"]
            assert value["verdict"].startswith("Not established")
            assert len(value["matches"]) == 1
            assert "Not established" in json.loads(row["summary_json"])["Potential Duplication"]
        assert db.execute("SELECT relationship FROM pairs").fetchone()[0] == "Structural similarity review"


def test_portfolio_findings_roll_up_evidence_backed_rationalization(tmp_path):
    from excel_xray.portfolio import _findings

    with connect(tmp_path) as db:
        first = db.execute("INSERT INTO files(source_path,file_name,scan_status) VALUES(?,?,?)",
                           ("/one.xlsx", "one.xlsx", "full")).lastrowid
        second = db.execute("INSERT INTO files(source_path,file_name,scan_status) VALUES(?,?,?)",
                            ("/two.xlsx", "two.xlsx", "full")).lastrowid
        first_assessment = {"file": {
            "potential_duplication": {"value": {"verdict": "Potential functional duplication — review",
                "matches": [{"comparison_id": str(second), "file": "two.xlsx"}]}},
            "similar_duplicate_files": {"value": [{"comparison_id": str(second), "file": "two.xlsx"}]},
            "potential_simplification": {"value": {"candidates": [{"worksheet": "Calc"}]}},
            "potential_consolidation": {"value": {"verdict": "Potential consolidation — review",
                "candidates": [{"comparison_id": str(second), "file": "two.xlsx"}]}},
            "potential_automation": {"value": {"candidates": [{
                "worksheet": "Input", "observed_manual_step": "Manual paste of data",
                "evidence": [{"text": "Manual input"}],
            }]}},
            "potential_retirement": {"value": {"verdict": "Potential retirement — owner review"}},
        }, "tabs": []}
        db.execute("UPDATE files SET assessment_json=? WHERE id=?",
                   (json.dumps(first_assessment), first))
        _findings(db)
        titles = {row[0] for row in db.execute("SELECT title FROM findings")}
        assert {
            "Potential functional duplication", "Similar EUC business use cases",
            "Potential simplification", "Potential consolidation", "Potential automation",
            "Potential retirement",
        } <= titles


def test_two_thousand_records_compare_without_individual_reports(tmp_path):
    with connect(tmp_path) as db:
        for i in range(2000):
            fp = {"file_id": str(i), "file_name": f"f{i}.xlsx", "logic_type": "Calculation",
                  "skeletons": [f"shape_{i}"], "headers": [f"header_{i}"],
                  "functions": [], "sheet_names": [f"sheet_{i}"]}
            db.execute("""INSERT INTO files(source_path,file_id,file_name,scan_status,
                       summary_json,assessment_json,fingerprint_json) VALUES(?,?,?,?,?,?,?)""",
                       (f"/synthetic/f{i}.xlsx", str(i), fp["file_name"], "full", "{}",
                        json.dumps({"file": {}, "tabs": [], "review": {}}), json.dumps(fp)))
        _compare(db)
        assert db.execute("SELECT COUNT(*) FROM files").fetchone()[0] == 2000
        assert db.execute("SELECT COUNT(*) FROM pairs").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM files WHERE assessment_json IS NOT NULL").fetchone()[0] == 2000


def test_identical_templates_keep_a_bounded_match_shortlist(tmp_path):
    with connect(tmp_path) as db:
        for i in range(100):
            fp = {"file_id": str(i), "file_name": f"f{i}.xlsx", "logic_type": "Calculation",
                  "skeletons": ["same_shape"], "headers": ["same_header"],
                  "functions": ["SUM"], "sheet_names": ["same_sheet"]}
            db.execute("""INSERT INTO files(source_path,file_id,file_name,scan_status,
                       summary_json,assessment_json,fingerprint_json) VALUES(?,?,?,?,?,?,?)""",
                       (f"/synthetic/f{i}.xlsx", str(i), fp["file_name"], "full", "{}",
                        json.dumps({"file": {}, "tabs": [], "review": {}}), json.dumps(fp)))
        _compare(db)
        stored = db.execute("SELECT COUNT(*) FROM pairs").fetchone()[0]
        assert 0 < stored <= 100 * 20
        assert stored < 100 * 99 // 2
